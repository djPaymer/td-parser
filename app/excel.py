"""Excel input (manufacturer list) and output (products + per-site summary)."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.parsers.product import ProductData

URL_HEADER_RE = re.compile(r"url|site|website|web|link|domain|сайт|ссылк|домен|адрес", re.I)
NAME_HEADER_RE = re.compile(r"name|manufacturer|brand|vendor|company|производител|бренд|назван|компани|поставщик", re.I)
URL_LIKE_RE = re.compile(r"^(?:https?://)?(?:www\.)?[\w.-]+\.[a-z]{2,}(?:[/?#].*)?$", re.I)

EXCEL_CELL_LIMIT = 32_000  # Excel allows 32767 characters per cell

PRODUCT_COLUMNS = (
    "Производитель", "Сайт", "URL товара", "Название", "Артикул", "Описание", "Характеристики", "Изображение",
)
SUMMARY_COLUMNS = (
    "Производитель", "Сайт", "Статус", "Товаров", "Regex товара", "Источник", "Уверенность",
    "Страниц загружено", "Комментарий",
)


@dataclass(slots=True)
class Manufacturer:
    name: str
    url: str
    row: int


@dataclass(slots=True)
class SiteReport:
    manufacturer: str
    site: str
    status: str  # ok | empty | error
    products: list[ProductData]
    regex: str = ""
    source: str = ""
    confidence: float = 0.0
    pages: int = 0
    note: str = ""


# --------------------------------------------------------------------------- input


def read_manufacturers(path: str | Path) -> list[Manufacturer]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [[_cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()

    url_col, name_col, start = _detect_columns(rows)
    out: list[Manufacturer] = []
    seen: set[str] = set()
    for index, row in enumerate(rows[start:], start=start + 1):
        url = row[url_col] if url_col is not None and url_col < len(row) else ""
        name = row[name_col] if name_col is not None and name_col < len(row) else ""
        if url_col is None:
            url = next((c for c in row if c and URL_LIKE_RE.match(c)), "")
            name = next((c for c in row if c and c != url), "")
        url = normalize_site_url(url)
        if not url:
            continue
        host = urlparse(url).netloc.lower()
        if host in seen:
            continue
        seen.add(host)
        out.append(Manufacturer(name=name or host.removeprefix("www."), url=url, row=index))
    return out


def normalize_site_url(value: str) -> str:
    value = (value or "").strip().strip("/")
    if not value or not URL_LIKE_RE.match(value):
        return ""
    if not value.lower().startswith(("http://", "https://")):
        value = "https://" + value
    parsed = urlparse(value)
    return f"{parsed.scheme}://{parsed.netloc}/"


def _detect_columns(rows: list[list[str]]) -> tuple[int | None, int | None, int]:
    """(url column, name column, first data row) from a header in the first rows, if any."""

    for index, row in enumerate(rows[:5]):
        url_col = next((i for i, c in enumerate(row) if c and URL_HEADER_RE.search(c) and not URL_LIKE_RE.match(c)), None)
        if url_col is None:
            continue
        name_col = next((i for i, c in enumerate(row) if i != url_col and c and NAME_HEADER_RE.search(c)), None)
        if name_col is None:
            name_col = next((i for i, c in enumerate(row) if i != url_col and c), None)
        return url_col, name_col, index + 1
    return None, None, 0


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


# --------------------------------------------------------------------------- output


class ResultWorkbook:
    """Products and per-site summary; ``save`` may be called after every site."""

    def __init__(self, path: str | Path) -> None:
        self._path = Path(path)
        self._wb = Workbook()
        self._products = self._wb.active
        self._products.title = "Товары"
        self._summary = self._wb.create_sheet("Сводка")
        self._header(self._products, PRODUCT_COLUMNS, (22, 28, 60, 50, 18, 80, 80, 40))
        self._header(self._summary, SUMMARY_COLUMNS, (22, 28, 10, 10, 70, 12, 12, 16, 80))

    def add(self, report: SiteReport) -> None:
        for product in report.products:
            self._products.append([
                report.manufacturer,
                report.site,
                product.url,
                _fit(product.name),
                _fit(product.sku),
                _fit(product.description),
                _fit(product.specs_text),
                _fit(product.image),
            ])
        self._summary.append([
            report.manufacturer,
            report.site,
            report.status,
            len(report.products),
            report.regex,
            report.source,
            round(report.confidence, 2) if report.confidence else None,
            report.pages,
            _fit(report.note),
        ])

    def save(self) -> Path:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        self._wb.save(self._path)
        return self._path

    @staticmethod
    def _header(ws, columns: tuple[str, ...], widths: tuple[int, ...]) -> None:
        ws.append(list(columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"


def _fit(value: str) -> str:
    value = value or ""
    if len(value) > EXCEL_CELL_LIMIT:
        return value[: EXCEL_CELL_LIMIT - 1] + "…"
    # Excel rejects control characters other than tab / newline
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
